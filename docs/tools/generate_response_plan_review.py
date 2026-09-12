from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
DOC_PATH = ROOT / "docs" / "福建公路应急事件分类与处置预案草案.docx"
PLAN_TEXT_PATH = ROOT / "road-agent-core" / "src" / "test" / "resources" / "dispatch" / "compact-response-plans.txt"

EVENT_NAMES = {
    "DT01": "崩塌（落石）",
    "DT02": "滑坡（坡体位移）",
    "DT03": "泥石流",
    "DT04": "沉陷与塌陷",
    "DT05": "水毁",
    "ET101": "拥堵",
    "ET102": "明火（火灾）",
    "ET103": "抛撒物",
    "ET104": "设备故障",
    "ET105": "占用应急车道",
    "ET106": "交通事故",
    "ET107": "异常停车",
    "ET108": "浓雾检测",
    "ET109": "路障",
    "ET110": "施工",
    "ET112": "道路积雪",
}


def chinese_font(run, name: str = "SimSun") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def read_facts(document: Document) -> dict[str, str]:
    facts: dict[str, str] = {}
    current_type: str | None = None
    awaiting_facts = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        match = re.search(r"\b(DT0[1-5]|ET10[1-9]|ET110|ET112)\b", text)
        if paragraph.style.name == "Heading 1" and match:
            current_type = match.group(1)
            awaiting_facts = False
        elif current_type and text in {"需要补充的现场信息", "需核实的现场信息"}:
            awaiting_facts = True
        elif current_type and awaiting_facts and text:
            facts[current_type] = text
            awaiting_facts = False
    missing = set(EVENT_NAMES) - set(facts)
    if missing:
        raise ValueError(f"原文档缺少现场事实项: {sorted(missing)}")
    return facts


def read_plans() -> dict[str, str]:
    plans: dict[str, str] = {}
    for line in PLAN_TEXT_PATH.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        event_type, text = line.split("\t", 1)
        plans[event_type] = text.replace("\\n", "\n").strip()
    missing = set(EVENT_NAMES) - set(plans)
    if missing:
        raise ValueError(f"简明预案缺少事件类型: {sorted(missing)}")
    return plans


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)

    title = document.styles["Title"]
    title.font.color.rgb = None
    title.font.size = Pt(24)
    title.font.bold = True
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.color.rgb = None
        style.font.name = "SimSun"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    if "Document Subtitle" not in [style.name for style in document.styles]:
        subtitle = document.styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = document.styles["Document Subtitle"]
    subtitle.font.name = "SimSun"
    subtitle.font.size = Pt(12)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")


def format_paragraph(paragraph, *, first_line: bool = True) -> None:
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.45
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(22)
    for run in paragraph.runs:
        chinese_font(run)
        run.font.size = Pt(11)


def build_document() -> None:
    document = Document(DOC_PATH)
    facts = read_facts(document)
    plans = read_plans()
    clear_body(document)
    configure_styles(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("福建公路应急事件处置预案审阅稿")
    chinese_font(run)

    subtitle = document.add_paragraph(style="Document Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("16类事件与智能体调用文本")
    chinese_font(run)

    intro = document.add_paragraph()
    intro.add_run("使用说明：").bold = True
    intro.add_run("智能体先根据事件类型选择对应预案，再结合现场事实和人工意见调整处置内容。本文件不设定资源类型或数量基线；资源需求应按事件实际情况提出，再由系统从已启用的数据库库存中匹配。未核实的事实不得臆造。")
    format_paragraph(intro, first_line=False)

    for index, event_type in enumerate(EVENT_NAMES, start=1):
        document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"{index}  {event_type}  {EVENT_NAMES[event_type]}")
        chinese_font(run)

        h2 = document.add_paragraph(style="Heading 2")
        h2.paragraph_format.keep_with_next = True
        run = h2.add_run("需核实的现场信息")
        chinese_font(run)
        paragraph = document.add_paragraph(facts[event_type])
        format_paragraph(paragraph)

        h2 = document.add_paragraph(style="Heading 2")
        h2.paragraph_format.keep_with_next = True
        run = h2.add_run("应急处置预案")
        chinese_font(run)
        for block in plans[event_type].split("\n\n"):
            paragraph = document.add_paragraph(block.strip())
            format_paragraph(paragraph)

    header = document.sections[0].header.paragraphs[0]
    header.text = "福建公路应急事件处置预案"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        chinese_font(run)
        run.font.size = Pt(9)

    document.core_properties.title = "福建公路应急事件处置预案审阅稿"
    document.core_properties.subject = "16类事件的智能体调用预案，不限定应急资源基线"
    document.save(DOC_PATH)


if __name__ == "__main__":
    build_document()
